import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

ORANGE = RGBColor(0xFF, 0x66, 0x00)
NAVY = RGBColor(0x1B, 0x3F, 0x8B)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x55, 0x55)
TINT = RGBColor(0xFF, 0xF0, 0xE5)    # light orange
NTINT = RGBColor(0xEA, 0xF0, 0xFA)   # light navy
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xF5, 0xC8, 0xA8)
GREEN = RGBColor(0x1E, 0x8E, 0x3E)
RED = RGBColor(0xC6, 0x28, 0x28)
FONT = "Times New Roman"

prs = Presentation("base.pptx")
slides = list(prs.slides)
_n = [0]


def nxt():
    s = slides[_n[0]]
    _n[0] += 1
    return s


def set_title(slide, text):
    tb = [sh for sh in slide.shapes if sh.name == "TextBox 14"][0]
    # center the title over the orange band (band spans ~228pt..582pt)
    tb.left, tb.width = Pt(228), Pt(354)
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.text = text
    size = 24 if len(text) <= 24 else 22 if len(text) <= 27 else 20
    r.font.size = Pt(size)
    tb.text_frame.word_wrap = False


def _font(run, size, color=INK, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, space=4, bullet=False, margin=0.05):
    """paras: list of str or list of [(text, opts)] runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(margin))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", italic))
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.25)))
            pPr.set("indent", str(-Inches(0.25)))
            bc = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(bc, qn("a:srgbClr")).set("val", "FF6600")
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "▪")
    return tb


def box(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08,
        shadow=False, dash=False, lw=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = 4  # dash
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if not shadow:
        sp = s._element.spPr
        etree.SubElement(sp, qn("a:effectLst"))
    s.text_frame.text = ""
    return s


def shape_text(s, paras, size=14, color=INK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.08):
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(margin))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", False))


def badge(slide, x, y, d, label, fill=ORANGE, size=16):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    shape_text(c, [label], size=size, color=WHITE, bold=True, margin=0)
    return c


def arrow(slide, x1, y1, x2, y2, color=NAVY, w=1.75, dash=False, both=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if both:
        etree.SubElement(ln, qn("a:headEnd")).set("type", "triangle")
    etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")
    return c


def table(slide, x, y, w, rows, colw, size=12, rowh=0.4, header_fill=NAVY):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows))).table
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (TINT if i % 2 == 0 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            _font(r, size, WHITE if i == 0 else INK, bold=(i == 0 or j == 0))
    return t


# ---------------------------------------------------------------- 1. Title
s = nxt()
set_title(s, "API & Microservices Project")
text(s, 0.6, 1.85, 9.8, 0.8, ["Campus Food Pre-Order System"], size=40, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
text(s, 1.0, 2.75, 9.0, 0.9,
     ["A Spring Boot microservices platform for multi-restaurant ordering with Order ID & OTP pickup verification"],
     size=18, color=GREY, italic=True, align=PP_ALIGN.CENTER)
card = box(s, 1.5, 4.0, 8.0, 2.35, fill=TINT)
text(s, 1.8, 4.15, 3.6, 0.4, ["Submitted by"], size=18, color=ORANGE, bold=True)
text(s, 1.8, 4.6, 3.7, 1.9, ["[Student Name 1] – [Roll No.]", "[Student Name 2] – [Roll No.]",
                             "[Student Name 3] – [Roll No.]", "[Student Name 4] – [Roll No.]"], size=16, space=6)
text(s, 5.7, 4.15, 3.6, 0.4, ["Under the guidance of"], size=18, color=ORANGE, bold=True)
text(s, 5.7, 4.6, 3.6, 1.2, ["[Guide Name]", "[Designation], Dept. of [Department]"], size=16, space=6)
text(s, 0.6, 7.0, 9.8, 0.8, [[("Department of [Department]", {"bold": True})],
                             "Aditya University, Surampalem, Andhra Pradesh"],
     size=16, color=NAVY, align=PP_ALIGN.CENTER, space=2)

# ---------------------------------------------------------------- 2. Introduction
s = nxt()
set_title(s, "Introduction")
text(s, 0.6, 1.7, 5.4, 5.8, [
    "College canteens still take orders through phone calls and write them in a notebook.",
    "When food is handed over, the order is struck off by hand – slow, error-prone, and with no proof of who ordered.",
    "Every canteen on campus faces the same problem.",
    "Our project brings all campus restaurants onto one website: students order online, restaurants manage today's menu, and food is collected at the counter using a unique Order ID and OTP.",
    "The backend is built as Spring Boot microservices that communicate through REST APIs, with a React + Tailwind frontend.",
], size=17, bullet=True, space=12)
# right: today vs with our system
for col, (head, fill, hc, steps) in enumerate([
    ("Today", NTINT, NAVY, ["Phone call", "Write in notebook", "Strike off by hand"]),
    ("With our system", TINT, ORANGE, ["Order on website", "Get Order ID + OTP", "Verify & collect"]),
]):
    y0 = 1.75 + col * 2.95
    box(s, 6.4, y0, 4.0, 2.65, fill=fill)
    text(s, 6.6, y0 + 0.12, 3.6, 0.4, [head], size=18, color=hc, bold=True)
    for k, st in enumerate(steps):
        yy = y0 + 0.65 + k * 0.62
        badge(s, 6.65, yy, 0.42, str(k + 1), fill=hc, size=14)
        text(s, 7.2, yy - 0.02, 3.1, 0.46, [st], size=16, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 3. Abstract
s = nxt()
set_title(s, "Abstract")
box(s, 0.6, 1.65, 9.8, 4.75, fill=TINT)
text(s, 0.9, 1.85, 9.2, 4.7, [
    "Most college canteens take orders over the phone and record them in a notebook, striking each one off after the food is handed over. "
    "This manual process leads to lost orders, long waiting times and disputes at the counter, and the same problem exists in every canteen on campus.",
    "This project presents the Campus Food Pre-Order System, a web application built on a microservices architecture using Spring Boot and REST APIs. "
    "Students choose from multiple campus restaurants, see only the dishes available that day, and place pickup orders. "
    "Every order generates a unique Order ID and a one-time password (OTP) shown only to that student. "
    "Restaurant staff verify the Order ID and OTP at the counter before handing over the food.",
    "Each restaurant manages only its own menu and orders, while a hidden administrator manages all restaurants, users and orders. "
    "The system replaces the notebook with a fast, secure and traceable digital workflow.",
], size=17, align=PP_ALIGN.JUSTIFY, space=12)
text(s, 0.6, 6.95, 1.4, 0.4, ["Keywords:"], size=16, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
kx = 1.95
for kw, w in [("Microservices", 1.45), ("Spring Boot", 1.3), ("REST API", 1.1), ("API Gateway", 1.35),
              ("Eureka", 0.9), ("JWT", 0.7), ("OTP", 0.7)]:
    c = box(s, kx, 6.97, w, 0.38, fill=NAVY, radius=0.5)
    shape_text(c, [kw], size=13, color=WHITE, bold=True, margin=0.02)
    kx += w + 0.12

# ---------------------------------------------------------------- 4. Problem statement
s = nxt()
set_title(s, "Problem Statement")
text(s, 0.6, 1.6, 9.8, 0.8, [
    "Campus canteens manage orders manually through phone calls and a notebook, which causes the following problems:"],
    size=17, color=INK)
probs = [
    ("Phone-based ordering", "Lines are busy at peak hours, so calls and orders are missed."),
    ("Handwritten records", "Notebook entries are hard to read, easy to mix up and easy to lose."),
    ("No order verification", "Anyone can claim an order – there is no proof of who ordered."),
    ("Unknown availability", "Students don't know which dishes are available today."),
    ("No records or reports", "No data on daily orders, sales or popular dishes."),
    ("Common to all canteens", "Every canteen on campus faces the very same problem."),
]
for i, (h, d) in enumerate(probs):
    cx = 0.6 + (i % 3) * 3.33
    cy = 2.45 + (i // 3) * 2.6
    box(s, cx, cy, 3.1, 2.35, fill=TINT)
    badge(s, cx + 0.2, cy + 0.2, 0.5, str(i + 1), size=16)
    text(s, cx + 0.2, cy + 0.82, 2.75, 0.45, [h], size=17, color=NAVY, bold=True)
    text(s, cx + 0.2, cy + 1.25, 2.75, 1.05, [d], size=15, color=INK)

# ---------------------------------------------------------------- 5. Existing vs Proposed
s = nxt()
set_title(s, "Existing vs Proposed System")
rows = [
    ("Orders taken by phone call", "Orders placed on a website"),
    ("Written in a notebook", "Stored safely in a database"),
    ("Struck off by hand", "Status updated: Placed → Ready → Collected"),
    ("No verification of the customer", "Order ID + OTP verified at the counter"),
    ("Menu availability unknown", "Only today's available dishes are shown"),
    ("No records or reports", "Admin reports on orders and sales"),
]
for col, (head, fill, hc, mark, mc) in enumerate([
    ("Existing System", NTINT, NAVY, "✗", RED),
    ("Proposed System", TINT, ORANGE, "✓", GREEN),
]):
    x0 = 0.6 + col * 5.0
    hb = box(s, x0, 1.65, 4.8, 0.6, fill=hc)
    shape_text(hb, [head], size=20, color=WHITE, bold=True)
    box(s, x0, 2.35, 4.8, 5.4, fill=fill)
    for k, r in enumerate(rows):
        yy = 2.55 + k * 0.85
        text(s, x0 + 0.2, yy, 0.4, 0.6, [mark], size=22, color=mc, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x0 + 0.65, yy, 4.0, 0.6, [r[col]], size=16, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 6. Objectives
s = nxt()
set_title(s, "Objectives")
objs = [
    "Replace the phone-and-notebook system with online ordering for all campus restaurants.",
    "Let students order from multiple restaurants and see only the dishes available today.",
    "Generate a unique Order ID and OTP for every order to ensure safe, verified pickup.",
    "Give each restaurant a dashboard to manage only its own menu, availability and orders.",
    "Provide a hidden administrator who can manage restaurants, users and all orders.",
    "Build the system as independent Spring Boot microservices communicating through REST APIs.",
]
for i, o in enumerate(objs):
    yy = 1.7 + i * 1.0
    box(s, 0.6, yy, 9.8, 0.82, fill=TINT if i % 2 == 0 else NTINT)
    badge(s, 0.8, yy + 0.16, 0.5, str(i + 1), fill=ORANGE if i % 2 == 0 else NAVY)
    text(s, 1.5, yy, 8.7, 0.82, [o], size=17, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 7. Technologies
s = nxt()
set_title(s, "Technologies Used")
techs = [
    ("Backend", ["Java 17", "Spring Boot 3", "Spring Initializr", "Spring Web – REST APIs"]),
    ("Microservices", ["Spring Cloud Gateway", "Netflix Eureka Server", "OpenFeign", "Resilience4j"]),
    ("Security", ["Spring Security", "JWT tokens", "Role-based access", "BCrypt hashing"]),
    ("Database", ["MySQL", "Spring Data JPA", "Hibernate", "Database per service"]),
    ("Frontend", ["React (Vite)", "Tailwind CSS", "Axios", "React Router"]),
    ("Tools", ["Swagger (OpenAPI)", "Postman", "Maven", "Git & GitHub"]),
]
for i, (h, items) in enumerate(techs):
    cx = 0.6 + (i % 3) * 3.33
    cy = 1.7 + (i // 3) * 3.05
    fill, hc = (TINT, ORANGE) if (i % 2 == 0) else (NTINT, NAVY)
    box(s, cx, cy, 3.1, 2.85, fill=fill)
    hb = box(s, cx, cy, 3.1, 0.6, fill=hc)
    shape_text(hb, [h], size=19, color=WHITE, bold=True)
    text(s, cx + 0.2, cy + 0.75, 2.8, 2.05, items, size=16, bullet=True, space=6)

# ---------------------------------------------------------------- 8. System architecture
s = nxt()
set_title(s, "System Architecture")
xs = [0.6, 3.3, 6.0]
W = 2.3
for x, (t1, t2) in zip(xs, [("Student Portal", "React · Tailwind"), ("Restaurant Portal", "React · Tailwind"),
                            ("Admin Portal", "React · hidden URL")]):
    b = box(s, x, 1.65, W, 0.75, fill=NTINT, line=NAVY)
    shape_text(b, [[(t1, {"bold": True, "color": NAVY})], [(t2, {"size": 12, "color": GREY})]], size=16)
    arrow(s, x + W / 2, 2.4, x + W / 2, 2.85)
g = box(s, 0.6, 2.85, 7.7, 0.85, fill=ORANGE)
shape_text(g, [[("API Gateway", {"bold": True, "size": 19}), ("   (Spring Cloud Gateway · port 8080)", {"size": 14})],
               [("Single entry point · routes requests · validates JWT & role", {"size": 13})]], color=WHITE)
svc = [("Auth Service", ":8081", "Login · JWT · roles", "auth_db"),
       ("Restaurant Service", ":8082", "Dishes · availability", "restaurant_db"),
       ("Order Service", ":8083", "Orders · ID + OTP", "order_db")]
for x, (n, port, d, db) in zip(xs, svc):
    arrow(s, x + W / 2, 3.7, x + W / 2, 4.15)
    b = box(s, x, 4.15, W, 1.05, fill=TINT, line=ORANGE)
    shape_text(b, [[(n, {"bold": True, "color": NAVY})], [("port " + port[1:], {"size": 12, "color": GREY})],
                   [(d, {"size": 13})]], size=15)
    arrow(s, x + W / 2, 5.2, x + W / 2, 5.6, both=True, w=1.25)
    c = box(s, x + 0.45, 5.6, 1.4, 1.0, fill=NAVY, shape=MSO_SHAPE.CAN)
    shape_text(c, [[("MySQL", {"size": 12})], [(db, {"bold": True, "size": 12})]], color=WHITE, margin=0.02)
# order -> restaurant REST call
arrow(s, 6.0, 4.68, 5.6, 4.68, color=ORANGE, w=2)
# eureka
e = box(s, 8.75, 2.85, 1.75, 2.35, fill=NTINT, line=NAVY, dash=True, lw=1.5)
shape_text(e, [[("Eureka Server", {"bold": True, "color": NAVY})], [(":8761", {"size": 13, "color": GREY})],
               [("Service registry – all services register & discover each other", {"size": 12})]], size=16)
arrow(s, 8.3, 3.27, 8.75, 3.27, dash=True, w=1.25, both=True)
arrow(s, 8.3, 4.68, 8.75, 4.68, dash=True, w=1.25, both=True)
text(s, 0.6, 6.95, 9.9, 0.8, [
    [("Flow: ", {"bold": True, "color": NAVY}),
     ("React app → API Gateway → microservice → its own MySQL database. Clients never call services directly; "
      "the Order Service calls the Restaurant Service via REST to check dishes and prices.", {})]],
    size=15)

# ---------------------------------------------------------------- 9. Modules
s = nxt()
set_title(s, "Modules")
mods = [
    ("Student", ORANGE, TINT, ["Register / login", "Choose a restaurant", "See only today's available dishes",
                               "Add to cart & place order", "Get Order ID + OTP on screen", "Track order status",
                               "Collect at counter (no room delivery)"]),
    ("Restaurant", NAVY, NTINT, ["Access to own restaurant only", "Add full dish list once",
                                 "Mark dishes available / unavailable daily", "Receive orders live",
                                 "Update status: Preparing → Ready", "Verify Order ID + OTP", "Mark order Collected"]),
    ("Admin (hidden)", ORANGE, TINT, ["Not visible to anyone", "Logs in via hidden URL only", "Add / remove restaurants",
                                      "Create restaurant logins", "Manage all dishes & users", "View all orders",
                                      "Reports: sales & popular dishes"]),
]
for i, (h, hc, fill, items) in enumerate(mods):
    cx = 0.6 + i * 3.33
    box(s, cx, 1.65, 3.1, 6.15, fill=fill)
    badge(s, cx + 1.15, 1.8, 0.8, h[0], fill=hc, size=26)
    text(s, cx + 0.1, 2.68, 2.9, 0.45, [h], size=20, color=hc, bold=True, align=PP_ALIGN.CENTER)
    text(s, cx + 0.2, 3.25, 2.8, 4.5, items, size=15, bullet=True, space=9)

# ---------------------------------------------------------------- 10. Implementation – services
s = nxt()
set_title(s, "Implementation – Services")
text(s, 0.6, 1.6, 9.8, 0.75, [
    "Each microservice is a separate Spring Boot project created with Spring Initializr and runs on its own port:"],
    size=17)
table(s, 0.6, 2.4, 9.8, [
    ["Microservice", "Port", "Responsibility", "Database"],
    ["Eureka Server", "8761", "Service registry and discovery", "–"],
    ["API Gateway", "8080", "Single entry point, routing, JWT and role check", "–"],
    ["Auth Service", "8081", "Register, login, JWT, roles (STUDENT / RESTAURANT / ADMIN)", "auth_db"],
    ["Restaurant Service", "8082", "Restaurants, dishes, daily availability", "restaurant_db"],
    ["Order Service", "8083", "Orders, Order ID & OTP generation, status, OTP verification", "order_db"],
], colw=[2.2, 0.9, 4.9, 1.8], size=15, rowh=0.56)
box(s, 0.6, 6.3, 9.8, 1.5, fill=NTINT)
text(s, 0.85, 6.38, 9.3, 1.38, [
    [("Database per service: ", {"bold": True, "color": NAVY}),
     ("every service owns its data and shares it only through REST APIs.", {})],
    [("Frontend: ", {"bold": True, "color": NAVY}),
     ("React (Vite) + Tailwind CSS app calls only the API Gateway using Axios.", {})],
    [("API docs: ", {"bold": True, "color": NAVY}),
     ("Swagger UI for every service; Resilience4j circuit breaker on service-to-service calls.", {})],
], size=16, space=8)

# ---------------------------------------------------------------- 11. Implementation – order flow
s = nxt()
set_title(s, "Implementation – Order Flow")
steps = [
    ("Student", "Logs in and selects a restaurant"),
    ("Student", "Sees only dishes available today"),
    ("Student", "Places order (pickup only)"),
    ("System", "Generates Order ID + 4-digit OTP"),
    ("Restaurant", "Gets the order, prepares it, marks Ready"),
    ("Student", "Comes to counter, tells Order ID + OTP"),
    ("Restaurant", "Enters Order ID + OTP in dashboard"),
    ("System", "Verifies OTP – if it matches, order is Collected"),
]
cw, chh, gap = 2.15, 2.25, 0.4
for i, (who, what) in enumerate(steps):
    row, k = divmod(i, 4)
    cx = 0.6 + k * (cw + gap)
    cy = 1.75 + row * 2.75
    fill = TINT if who != "Restaurant" else NTINT
    box(s, cx, cy, cw, chh, fill=fill)
    badge(s, cx + 0.15, cy + 0.15, 0.5, str(i + 1), fill=NAVY if who == "Restaurant" else ORANGE)
    text(s, cx + 0.72, cy + 0.17, 1.4, 0.45, [who], size=14, color=GREY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx + 0.15, cy + 0.8, cw - 0.3, 1.4, [what], size=16, bold=False)
    if k < 3:
        arrow(s, cx + cw + 0.05, cy + chh / 2, cx + cw + gap - 0.05, cy + chh / 2, color=ORANGE, w=2.25)
text(s, 0.6, 7.25, 9.8, 0.5, [
    [("Note: ", {"bold": True, "color": NAVY}),
     ("the OTP is shown only on the ordering student's screen, so only that student can collect the order.", {})]],
    size=15)

# ---------------------------------------------------------------- 12. Implementation – REST APIs
s = nxt()
set_title(s, "Implementation – REST APIs")
table(s, 0.6, 1.65, 9.8, [
    ["Method", "Endpoint", "Access", "Purpose"],
    ["POST", "/api/auth/register", "Public", "Student registration"],
    ["POST", "/api/auth/login", "Public", "Login, returns JWT token"],
    ["GET", "/api/restaurants", "Student", "List campus restaurants"],
    ["GET", "/api/restaurants/{id}/dishes", "Student", "Only today's available dishes"],
    ["POST", "/api/orders", "Student", "Place order → Order ID + OTP"],
    ["GET", "/api/orders/my", "Student", "My orders and their status"],
    ["POST", "/api/restaurant/dishes", "Restaurant", "Add a dish to own menu"],
    ["PATCH", "/api/restaurant/dishes/{id}/availability", "Restaurant", "Mark available / unavailable"],
    ["GET", "/api/restaurant/orders", "Restaurant", "Own restaurant's orders"],
    ["POST", "/api/restaurant/orders/verify", "Restaurant", "Verify Order ID + OTP → Collected"],
    ["ALL", "/api/admin/**", "Admin", "Manage everything (hidden)"],
], colw=[0.95, 4.1, 1.45, 3.3], size=13, rowh=0.5)

# ---------------------------------------------------------------- 13. Security & access control
s = nxt()
set_title(s, "Security & Access Control")
rules = [
    ("Role-based access", "Every request carries a JWT with the user's role, checked by the gateway and services."),
    ("Restaurant isolation", "A restaurant login is linked to one restaurant and can only see its own menu and orders."),
    ("Hidden admin", "No admin link, button or signup anywhere. Admin uses a hidden URL; accounts are created in the DB."),
    ("Students only order", "Students cannot open restaurant or admin pages or APIs (403 Forbidden)."),
    ("Only available dishes", "Unavailable dishes are filtered out by the API and never sent to students."),
    ("Pickup only", "No room delivery – food is collected at the counter with Order ID + OTP."),
]
for i, (h, d) in enumerate(rules):
    cx = 0.6 + (i % 2) * 5.0
    cy = 1.7 + (i // 2) * 2.0
    fill, hc = (TINT, ORANGE) if (i // 2 + i % 2) % 2 == 0 else (NTINT, NAVY)
    box(s, cx, cy, 4.8, 1.8, fill=fill)
    badge(s, cx + 0.2, cy + 0.22, 0.55, str(i + 1), fill=hc)
    text(s, cx + 0.95, cy + 0.2, 3.7, 0.5, [h], size=18, color=hc, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx + 0.95, cy + 0.72, 3.7, 1.05, [d], size=15)

# ---------------------------------------------------------------- 14. Future scope
s = nxt()
set_title(s, "Future Scope")
fut = [
    ("Online payments", "UPI / card payment while ordering."),
    ("Live notifications", "SMS or push alert when the order is ready."),
    ("Mobile app", "Android and iOS apps using the same REST APIs."),
    ("Scheduled orders", "Pre-order food for a chosen pickup time slot."),
    ("Ratings & feedback", "Students rate dishes and restaurants."),
    ("Cloud deployment", "Docker containers and cloud hosting for scale."),
]
for i, (h, d) in enumerate(fut):
    cx = 0.6 + (i % 3) * 3.33
    cy = 1.75 + (i // 3) * 3.0
    fill, hc = (TINT, ORANGE) if i % 2 == 0 else (NTINT, NAVY)
    box(s, cx, cy, 3.1, 2.7, fill=fill)
    badge(s, cx + 1.2, cy + 0.25, 0.7, str(i + 1), fill=hc, size=20)
    text(s, cx + 0.15, cy + 1.1, 2.8, 0.45, [h], size=18, color=hc, bold=True, align=PP_ALIGN.CENTER)
    text(s, cx + 0.2, cy + 1.6, 2.7, 1.0, [d], size=15, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 15. Conclusion
s = nxt()
set_title(s, "Conclusion")
for i, (num, lab) in enumerate([("3", "user roles"), ("5", "microservices"), ("1", "OTP per order")]):
    cx = 0.6 + i * 3.33
    hc = ORANGE if i != 1 else NAVY
    box(s, cx, 1.7, 3.1, 1.6, fill=TINT if i != 1 else NTINT)
    text(s, cx, 1.75, 3.1, 0.95, [num], size=48, color=hc, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx, 2.7, 3.1, 0.45, [lab], size=17, color=INK, align=PP_ALIGN.CENTER)
text(s, 0.6, 3.65, 9.8, 4.1, [
    "The Campus Food Pre-Order System replaces the phone-and-notebook process with a simple online workflow.",
    "Order ID + OTP verification makes sure the right student receives the right order.",
    "Restaurants manage only their own menu and daily availability, while a hidden admin oversees everything.",
    "Built as independent Spring Boot microservices with REST APIs, the system is easy to extend, scale and maintain – and can be used by every canteen on campus.",
], size=18, bullet=True, space=14)

# ---------------------------------------------------------------- 16. Thank you
s = nxt()
set_title(s, "Thank You")
text(s, 0.6, 2.6, 9.8, 1.2, ["Thank You!"], size=60, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.6, 3.9, 9.8, 0.6, ["Questions?"], size=26, color=ORANGE, italic=True, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.2, 9.8, 0.9, ["Campus Food Pre-Order System",
                             "API & Microservices · Aditya University, Surampalem"],
     size=16, color=GREY, align=PP_ALIGN.CENTER, space=4)

assert _n[0] == len(slides)
prs.save("Campus_Food_PreOrder_System.pptx")
print("saved")
